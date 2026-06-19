import os
import re
from flask import Flask, render_template, request, jsonify
import google.generativeai as genai
from dotenv import load_dotenv
import PyPDF2
import docx
import requests
from bs4 import BeautifulSoup

load_dotenv()

app = Flask(__name__)

# Configure Gemini API
API_KEY = os.getenv("GEMINI_API_KEY")
if API_KEY:
    genai.configure(api_key=API_KEY)
else:
    print("Warning: GEMINI_API_KEY not found in environment variables.")

def extract_text_from_pdf(file_stream):
    try:
        reader = PyPDF2.PdfReader(file_stream)
        text = ""
        for page in reader.pages:
            text += page.extract_text() + "\n"
        return text
    except Exception as e:
        raise Exception(f"Error reading PDF: {str(e)}")

def extract_text_from_docx(file_stream):
    try:
        doc = docx.Document(file_stream)
        return "\n".join([para.text for para in doc.paragraphs])
    except Exception as e:
        raise Exception(f"Error reading DOCX: {str(e)}")

def extract_text_from_url(url):
    try:
        # Some sites block default user agents, adding a standard one
        headers = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64)'}
        response = requests.get(url, headers=headers, timeout=10)
        response.raise_for_status()
        soup = BeautifulSoup(response.text, 'html.parser')
        
        # Remove script and style elements
        for script in soup(["script", "style"]):
            script.extract()
            
        text = soup.get_text(separator=' ')
        # Collapse whitespace
        lines = (line.strip() for line in text.splitlines())
        chunks = (phrase.strip() for line in lines for phrase in line.split("  "))
        text = '\n'.join(chunk for chunk in chunks if chunk)
        return text
    except Exception as e:
        raise Exception(f"Error fetching URL: {str(e)}")

def is_valid_url(text):
    # Basic URL validation
    regex = re.compile(
        r'^(?:http|ftp)s?://' # http:// or https://
        r'(?:(?:[A-Z0-9](?:[A-Z0-9-]{0,61}[A-Z0-9])?\.)+(?:[A-Z]{2,6}\.?|[A-Z0-9-]{2,}\.?)|' #domain...
        r'localhost|' #localhost...
        r'\d{1,3}\.\d{1,3}\.\d{1,3}\.\d{1,3})' # ...or ip
        r'(?::\d+)?' # optional port
        r'(?:/?|[/?]\S+)$', re.IGNORECASE)
    return re.match(regex, text.strip()) is not None

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/generate', methods=['POST'])
def generate():
    resume_text = ""
    job_description_text = ""
    tone = "Professional"
    output_type = "Standard Cold Email"
    cta = "Default"
    
    if request.form:
        resume_text = request.form.get('resume', '').strip()
        job_description_input = request.form.get('job_description', '').strip()
        tone = request.form.get('tone', 'Professional')
        output_type = request.form.get('output_type', 'Standard Cold Email')
        cta = request.form.get('cta', 'Default')
        
        # Check for file upload
        if 'resume_file' in request.files:
            file = request.files['resume_file']
            if file.filename != '':
                if file.filename.endswith('.pdf'):
                    resume_text = extract_text_from_pdf(file.stream)
                elif file.filename.endswith('.docx'):
                    resume_text = extract_text_from_docx(file.stream)
                else:
                    return jsonify({'error': 'Unsupported file format. Please upload a PDF or DOCX file.'}), 400
    else:
        return jsonify({'error': 'Invalid request format.'}), 400

    # Process Job Description (check if URL)
    if job_description_input:
        if is_valid_url(job_description_input):
            try:
                job_description_text = extract_text_from_url(job_description_input)
            except Exception as e:
                return jsonify({'error': str(e)}), 400
        else:
            job_description_text = job_description_input

    if not resume_text or not job_description_text:
        return jsonify({'error': 'Resume and Job Description are required.'}), 400
        
    try:
        model = genai.GenerativeModel('gemini-2.5-flash')
        
        # Adjust prompt based on output_type
        length_instruction = ""
        if "LinkedIn" in output_type:
            length_instruction = "This should be a short LinkedIn connection request (strictly around 300 characters). Keep it very brief and impactful."
        elif "Cover Letter" in output_type:
            length_instruction = "This should be a traditional, formal cover letter (about a full page). Follow standard letter formatting."
        else:
            length_instruction = "This should be a standard cold email (around 150-200 words). Make it punchy and readable."

        # Adjust CTA
        cta_instruction = ""
        if cta != "Default":
            cta_instruction = f"At the end of the message, include this specific Call-to-Action: '{cta}'."

        prompt = f"""
You are an expert career coach and copywriter. Your task is to write a highly personalized, compelling outreach message.
Using the provided resume and job description, create an outreach message that matches the specified tone and format.

FORMAT: {output_type}
TONE: {tone}

{length_instruction}
{cta_instruction}

JOB DESCRIPTION:
{job_description_text}

CANDIDATE'S RESUME:
{resume_text}

CRITICAL SYSTEM INSTRUCTIONS:
1. NEVER hallucinate or invent skills, experiences, degrees, or technologies. You must ONLY mention skills explicitly listed in the CANDIDATE'S RESUME.
2. If there is a missing piece of information (e.g., you don't know the Hiring Manager's Name or the Company Name), you MUST use a placeholder in brackets, exactly like [Company Name] or [Hiring Manager]. Do not guess or invent names.
3. Write a compelling subject line (unless it's a LinkedIn request).
4. The message should directly address how the candidate's actual skills solve the problems outlined in the job description.
5. Format the output clearly.
"""
        
        response = model.generate_content(prompt)
        
        return jsonify({'result': response.text})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/tweak', methods=['POST'])
def tweak():
    data = request.json
    original_text = data.get('original_text')
    instruction = data.get('instruction')
    
    if not original_text or not instruction:
        return jsonify({'error': 'Original text and instruction are required.'}), 400
        
    try:
        model = genai.GenerativeModel('gemini-2.5-flash')
        
        prompt = f"""
You are an expert editor. Your task is to rewrite the following outreach message according to the specific instruction below.

INSTRUCTION: {instruction}

CRITICAL SYSTEM INSTRUCTIONS:
1. NEVER hallucinate or invent skills, experiences, degrees, or technologies not present in the original message.
2. Retain any existing bracketed placeholders like [Company Name] or use them if information is missing.

ORIGINAL MESSAGE:
{original_text}
"""
        response = model.generate_content(prompt)
        return jsonify({'result': response.text})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

if __name__ == '__main__':
    app.run(debug=True, port=5000)
